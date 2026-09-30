import io
import sqlite3
import subprocess
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import Mock, patch

import fitz
from app import media_storage
from app.api import publishing
from app.auth import storage as auth_storage
from app.auth.security import create_access_token
from app.engines.publishing import (
    document_copy,
    material_copy,
    project_materials,
    storage,
)
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from scripts import migrate_sqlite_to_postgres


class MaterialCopyTests(unittest.TestCase):
    def setUp(self):
        directory = self.enterContext(
            tempfile.TemporaryDirectory(prefix=".copy-", dir=Path(__file__).parent),
        )
        self.db_path = str(Path(directory) / "test.db")
        self.media_root = Path(directory) / "media"
        self.enterContext(patch.object(auth_storage, "DB_PATH", self.db_path))
        self.enterContext(patch.object(storage, "DB_PATH", self.db_path))
        self.enterContext(patch.object(media_storage, "MEDIA_ROOT", str(self.media_root)))
        self.enterContext(patch.object(media_storage, "MEDIA_STORAGE_BACKEND", "local"))
        self.owner = auth_storage.create_user("owner", "owner@example.com", "test-hash")
        self.member = auth_storage.create_user("member", "member@example.com", "test-hash")
        self.admin = auth_storage.create_user("admin", "admin@example.com", "test-hash")
        self.outsider = auth_storage.create_user("outsider", "outsider@example.com", "test-hash")
        organization = auth_storage.get_current_organization(self.owner["id"])
        self.project = storage.create_manual_project(self.owner["id"], title="Copy project")
        with sqlite3.connect(self.db_path) as conn:
            for user, role in ((self.member, "member"), (self.admin, "admin")):
                conn.execute(
                    "INSERT INTO organization_memberships VALUES (?, ?, 'member', '2026-01-01')",
                    (organization["id"], user["id"]),
                )
                conn.execute(
                    "UPDATE user_organization_preferences SET organization_id = ? WHERE user_id = ?",
                    (organization["id"], user["id"]),
                )
                conn.execute(
                    "INSERT INTO project_memberships VALUES (?, ?, ?, '2026-01-01')",
                    (self.project.id, user["id"], role),
                )
        app = FastAPI()
        app.include_router(publishing.router)
        self.client = self.enterContext(TestClient(app))
        self.base = f"/api/v1/publishing/projects/{self.project.id}"
        self.material_set = project_materials.create_project_material_set(
            self.owner["id"], self.project.id, name="Copy",
        )

    def headers(self, user=None):
        return {"Authorization": "Bearer " + create_access_token((user or self.owner)["id"])}

    def create_copy(self, **changes):
        return self.client.post(
            self.base + "/materials/copy", headers=self.headers(),
            json={
                "material_set_id": self.material_set.id,
                "title": "Campaign.v1",
                "content": "<p>Hello <strong>world</strong></p>",
                **changes,
            },
        )

    def upload(self, filename, content=b"copy", mime="text/html", material_set_id=None):
        return self.client.post(
            self.base + "/materials", headers=self.headers(),
            data={"material_set_id": material_set_id or self.material_set.id},
            files={"file": (filename, content, mime)},
        )

    def preview(self, item, user=None):
        return self.client.get(
            self.base + f"/materials/{item['id']}/content", headers=self.headers(user),
        )

    def test_material_plain_text_import_preserves_rich_source_and_access(self):
        created = self.create_copy(content="<h2>Launch 😀</h2><p>Hello <strong>world</strong><br>Next line</p>")
        self.assertEqual(created.status_code, 200, created.text)
        item = created.json()["data"]
        original = self.preview(item).json()["data"]
        path = self.base + f"/materials/{item['id']}/content"
        plain = self.client.get(path, headers=self.headers(), params={"format": "text"})
        self.assertEqual(plain.status_code, 200, plain.text)
        self.assertEqual(plain.json()["data"], {
            "format": "text", "content": "Launch 😀\n\nHello world\nNext line",
        })
        self.assertEqual(self.preview(item).json()["data"], original)
        self.assertEqual(self.client.get(
            path, headers=self.headers(), params={"format": "invalid"},
        ).status_code, 422)
        self.assertEqual(self.client.get(path, params={"format": "text"}).status_code, 401)
        self.assertEqual(self.client.get(
            path, headers=self.headers(self.outsider), params={"format": "text"},
        ).status_code, 404)

    def edit(self, item, user=None, **changes):
        return self.client.patch(
            self.base + f"/materials/{item['id']}/content", headers=self.headers(user),
            json={"content": "<p>Edited</p>", **changes},
        )

    @staticmethod
    def pdf_bytes(text="PDF body", pages=1):
        with fitz.open() as document:
            for _ in range(pages):
                page = document.new_page()
                if text:
                    page.insert_text((72, 72), text)
            return document.tobytes()

    @staticmethod
    def docx_bytes():
        document = Document()
        document.add_heading("Document title", 1)
        document.add_paragraph().add_run("Bold body").bold = True
        document.add_paragraph("List item", "List Bullet")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table cell"
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def legacy(self, extension, data, mime=None):
        key = f"project-materials/legacy/{len(list(self.media_root.rglob('*'))) if self.media_root.exists() else 0}{extension}"
        media_storage.put_media_bytes(key, data)
        return project_materials.create_project_material(
            self.owner["id"], self.project.id, name="Legacy.v1" + extension,
            media_type="document", mime_type=mime or material_copy.DOCUMENT_CONTENT_TYPES.get(extension, "text/html"),
            file_size=len(data), object_key=key, material_set_id=self.material_set.id,
        ).model_dump()

    def test_upload_document_types_parse_to_database_html_without_media_writes(self):
        sources = {
            ".txt": ("原文 <script>not executed</script>\nSecond line".encode(), "&lt;script&gt;"),
            ".md": (b"# Title\n\n**bold** and *emphasis*", "<h2>Title</h2>"),
            ".markdown": (b"- Item\n\n> Quote", "<blockquote>"),
            ".pdf": (self.pdf_bytes(), "PDF body"),
            ".docx": (self.docx_bytes(), "<strong>Bold body</strong>"),
        }
        for extension, (source, expected) in sources.items():
            with self.subTest(extension=extension):
                with patch.object(publishing, "put_media_bytes") as put:
                    response = self.upload("release.v1" + extension.upper(), source)
                put.assert_not_called()
                self.assertEqual(response.status_code, 200, response.text)
                item = response.json()["data"]
                self.assertEqual(item["media_type"], "document")
                self.assertEqual(item["mime_type"], "text/html")
                self.assertTrue(item["name"].startswith("release.v1"))
                self.assertEqual(item["file_url"], "")
                self.assertEqual(item["object_key"], "")
                self.assertNotIn("content_html", item)
                preview = self.preview(item)
                self.assertEqual(preview.status_code, 200, preview.text)
                result = preview.json()["data"]
                self.assertEqual(result["format"], "html")
                self.assertIn(expected, result["content"])
                self.assertEqual(item["file_size"], len(result["content"].encode()))
                self.assertEqual(self.edit(item).status_code, 200)
                self.assertEqual(self.preview(item).json()["data"]["content"], "<p>Edited</p>")
        self.assertFalse(self.media_root.exists())
        for extension in (".html", ".htm", ".doc", ".exe", ""):
            self.assertEqual(self.upload("blocked" + extension).status_code, 400)
        for name, mime in (("photo.png", "image/png"), ("clip.mp4", "video/mp4")):
            item = self.upload(name).json()["data"]
            self.assertEqual(item["mime_type"], mime)

    def test_rich_create_preview_list_unique_names_rename_and_delete(self):
        content = '<h2>Title</h2><p>你好 <strong>world</strong><br><a href="https://example.com">Link</a></p>'
        first = self.create_copy(content=content)
        self.assertEqual(first.status_code, 200, first.text)
        item = first.json()["data"]
        second = self.create_copy().json()["data"]
        self.assertEqual(item["name"], "Campaign.v1")
        self.assertEqual(second["name"], "Campaign.v1 (1)")
        self.assertEqual(item["mime_type"], "text/html")
        self.assertEqual(item["media_type"], "document")
        self.assertEqual(item["file_size"], len(content.encode()))
        self.assertEqual(item["file_url"], "")
        self.assertEqual(item["object_key"], "")
        with sqlite3.connect(self.db_path) as conn:
            self.assertEqual(conn.execute(
                "SELECT content_html FROM project_materials WHERE id = ?", (item["id"],),
            ).fetchone()[0], content)
        response = self.preview(item)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["data"], {"content": content, "format": "html"})
        storage.init_db()
        storage.init_db()
        listed = self.client.get(
            self.base + "/materials", headers=self.headers(),
            params={"material_set_id": self.material_set.id},
        ).json()["data"]
        self.assertEqual({entry["name"] for entry in listed}, {"Campaign.v1", "Campaign.v1 (1)"})
        self.assertTrue(all(entry["file_url"] == "" for entry in listed))
        renamed = self.client.put(
            self.base + f"/materials/{item['id']}", headers=self.headers(),
            json={"name": "Renamed.v2.html"},
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["data"]["document_count"], 0)
        storage.init_db()
        self.assertEqual(
            project_materials.get_project_material(
                self.owner["id"], self.project.id, item["id"],
            ).name,
            "Renamed.v2.html",
        )
        self.assertEqual(self.client.delete(
            self.base + f"/materials/{item['id']}", headers=self.headers(),
        ).status_code, 200)
        self.assertFalse(self.media_root.exists())
        self.assertEqual(self.preview(item).status_code, 404)
        self.assertEqual(self.client.delete(
            self.base + f"/materials/{self.material_set.id}", headers=self.headers(),
        ).status_code, 200)
        self.assertEqual(self.preview(second).status_code, 404)

    def test_document_counts_never_enter_image_video_covers(self):
        self.assertEqual(self.material_set.document_count, 0)
        self.create_copy()
        self.upload("notes.md")
        photo = self.upload("photo.jpg").json()["data"]
        video = self.upload("clip.mp4").json()["data"]
        summary = self.client.get(
            self.base + "/materials", headers=self.headers(),
        ).json()["data"][0]
        self.assertEqual(
            [summary[key] for key in ("material_count", "image_count", "video_count", "document_count")],
            [4, 1, 1, 2],
        )
        self.assertEqual({cover["id"] for cover in summary["covers"]}, {photo["id"], video["id"]})
        renamed = self.client.put(
            self.base + f"/material-sets/{self.material_set.id}",
            headers=self.headers(), json={"name": "Renamed"},
        ).json()["data"]
        self.assertEqual(renamed["document_count"], 2)
        self.assertEqual(renamed["material_count"], 4)

    def test_copy_authorization_and_creator_manager_permissions(self):
        payload = {"material_set_id": self.material_set.id, "title": "Member", "content": "<p>copy</p>"}
        self.assertEqual(self.client.post(
            self.base + "/materials/copy", json=payload,
        ).status_code, 401)
        self.assertEqual(self.client.post(
            self.base + "/materials/copy", json=payload, headers=self.headers(self.outsider),
        ).status_code, 404)
        owner_item = self.create_copy().json()["data"]
        content_path = self.base + f"/materials/{owner_item['id']}/content"
        self.assertEqual(self.client.get(content_path).status_code, 401)
        self.assertEqual(self.preview(owner_item, self.outsider).status_code, 404)
        self.assertEqual(self.preview(owner_item, self.member).status_code, 200)
        for method, body in (("PUT", {"name": "Forbidden"}), ("DELETE", None)):
            response = self.client.request(
                method, self.base + f"/materials/{owner_item['id']}",
                headers=self.headers(self.member), json=body,
            )
            self.assertEqual(response.status_code, 403, response.text)
        member_response = self.client.post(
            self.base + "/materials/copy", json=payload, headers=self.headers(self.member),
        )
        self.assertEqual(member_response.status_code, 200, member_response.text)
        member_item = member_response.json()["data"]
        self.assertEqual(self.client.put(
            self.base + f"/materials/{member_item['id']}", headers=self.headers(self.member),
            json={"name": "Own edit"},
        ).status_code, 200)
        self.assertEqual(self.client.delete(
            self.base + f"/materials/{member_item['id']}", headers=self.headers(self.member),
        ).status_code, 200)
        self.assertEqual(self.client.delete(
            self.base + f"/materials/{owner_item['id']}", headers=self.headers(self.admin),
        ).status_code, 200)

    def test_project_and_set_validation_precedes_object_write(self):
        other_project = storage.create_manual_project(self.owner["id"], title="Other")
        other_set = project_materials.create_project_material_set(
            self.owner["id"], other_project.id, name="Other",
        )
        item = self.create_copy().json()["data"]
        with patch.object(publishing, "put_media_bytes") as put:
            for invalid_set in ("missing", other_set.id, item["id"], " "):
                self.assertEqual(self.create_copy(material_set_id=invalid_set).status_code, 404)
                self.assertEqual(self.upload("notes.txt", material_set_id=invalid_set).status_code, 404)
            put.assert_not_called()
        wrong_project_url = f"/api/v1/publishing/projects/{other_project.id}/materials/{item['id']}/content"
        self.assertEqual(self.client.get(wrong_project_url, headers=self.headers()).status_code, 404)
        self.assertEqual(self.preview({"id": self.material_set.id}).status_code, 404)
        self.assertEqual(self.client.get(
            self.base + "/materials", headers=self.headers(),
            params={"material_set_id": "missing"},
        ).status_code, 404)

    def test_sanitizer_drops_unsafe_content_attributes_and_urls(self):
        source = (
            '<script>alert(1)</script><style>bad</style><svg><script>bad</script></svg>'
            '<p onclick="bad()" style="color:red">Safe<img src=x onerror=bad()>'
            '<a href="jav&#x61;script:bad()" target="_blank">bad link</a>'
            '<a href="data:text/html,bad">data</a>'
            '<a href="//example.com">relative</a>'
            '<a href="https://example.com?q=&quot;x&quot;&amp;y=1" onclick="bad()">HTTPS</a>'
            '<a href="mailto:test@example.com">mail</a></p>'
        )
        response = self.create_copy(content=source)
        self.assertEqual(response.status_code, 200, response.text)
        cleaned = self.preview(response.json()["data"]).json()["data"]["content"]
        self.assertEqual(cleaned, (
            '<p>Safe<a>bad link</a><a>data</a><a>relative</a>'
            '<a href="https://example.com?q=&quot;x&quot;&amp;y=1">HTTPS</a>'
            '<a href="mailto:test@example.com">mail</a></p>'
        ))
        for href in ("javascript:bad", "java&#10;script:bad", "data:x", "file:///x", "https://", "/x"):
            self.assertEqual(
                material_copy.sanitize_copy_html(f'<a href="{href}">safe</a>'), "<a>safe</a>",
            )
        allowed = "<h2>H2</h2><h3>H3</h3><p><strong>S</strong><em>E</em><s>S</s><u>U</u><br></p><ul><li>U</li></ul><ol><li>O</li></ol><blockquote>B</blockquote>"
        self.assertEqual(material_copy.sanitize_copy_html(allowed), allowed)
        self.assertEqual(material_copy.sanitize_copy_html("<p><strong>x</p>"), "<p><strong>x</strong></p>")
        self.assertEqual(material_copy.sanitize_copy_html("a &lt;b&gt; & c"), "a &lt;b&gt; &amp; c")
        malformed = self.create_copy(content="<![unsupported]>text")
        self.assertIn(malformed.status_code, (200, 400))
        if malformed.status_code == 200:
            self.assertEqual(self.preview(malformed.json()["data"]).json()["data"]["content"], "text")

    def test_empty_and_oversize_copy_and_upload_rejected(self):
        for content in ("", "  ", "<p><br>&nbsp;</p>", "<script>bad</script>", "<img src=x>"):
            self.assertIn(self.create_copy(content=content).status_code, (400, 422))
        for title in ("", "  ", "a/b", "a\\b", "a" * 256):
            self.assertIn(self.create_copy(title=title).status_code, (400, 422))
        with patch.object(material_copy, "MAX_COPY_CONTENT_BYTES", 32):
            self.assertEqual(self.create_copy(content="中" * 11).status_code, 400)
            self.assertEqual(self.create_copy(content="&" * 10).status_code, 400)
        self.assertEqual(self.create_copy(content="a" * (1024 * 1024 + 1)).status_code, 422)
        self.assertEqual(self.upload("empty.txt", b"").status_code, 400)
        with patch.object(publishing, "MAX_UPLOAD_SIZE_BYTES", 4):
            self.assertEqual(self.upload("large.txt", b"12345").status_code, 413)
        self.assertFalse(self.media_root.exists())

    def test_missing_or_non_utf8_text_preview_errors(self):
        self.assertEqual(self.upload("binary.txt", b"\xff").status_code, 400)
        item = self.legacy(".txt", b"\xff")
        self.assertEqual(self.preview(item).status_code, 400)
        media_storage.delete_media(item["object_key"])
        self.assertEqual(self.preview(item).status_code, 404)
        item = self.upload("bom.txt", b"\xef\xbb\xbfHello").json()["data"]
        self.assertEqual(self.preview(item).json()["data"]["content"], "<p>Hello</p>")

    def test_failed_insert_cleans_up_copy_object(self):
        with patch.object(publishing, "create_project_material", side_effect=LookupError("Gone")):
            self.assertEqual(self.create_copy().status_code, 404)
        self.assertEqual(list(self.media_root.rglob("*.html")), [])

    def test_s3_legacy_read_preserved_and_new_copy_never_writes_objects(self):
        s3 = Mock()
        stored = {}
        s3.put_object.side_effect = lambda **kwargs: stored.update({kwargs["Key"]: kwargs})
        s3.get_object.side_effect = lambda **kwargs: {"Body": io.BytesIO(stored[kwargs["Key"]]["Body"])}
        s3.delete_object.side_effect = lambda **kwargs: stored.pop(kwargs["Key"])
        with (
            patch.object(media_storage, "MEDIA_STORAGE_BACKEND", "s3"),
            patch.object(media_storage, "MEDIA_S3_PREFIX", ""),
            patch.object(media_storage, "_s3_client", return_value=s3),
        ):
            legacy = self.legacy(".md", b"# Legacy")
            self.assertIn("<h2>Legacy</h2>", self.preview(legacy).json()["data"]["content"])
            self.assertEqual(self.edit(legacy).status_code, 200)
            self.assertEqual(stored[legacy["object_key"]]["Body"], b"# Legacy")
            self.assertEqual(self.preview(legacy).json()["data"]["content"], "<p>Edited</p>")
            s3.put_object.reset_mock()
            item = self.create_copy().json()["data"]
            self.assertEqual(self.preview(item).json()["data"]["format"], "html")
            document = self.upload("source.markdown").json()["data"]
            self.assertEqual(self.preview(document).json()["data"], {"format": "html", "content": "<p>copy</p>\n"})
            s3.put_object.assert_not_called()
            for entry in (item, document, legacy):
                self.assertEqual(self.client.delete(
                    self.base + f"/materials/{entry['id']}", headers=self.headers(),
                ).status_code, 200)
            self.assertEqual(stored, {})

    def test_patch_permissions_validation_and_title_preservation(self):
        item = self.create_copy().json()["data"]
        original = self.preview(item).json()["data"]
        path = self.base + f"/materials/{item['id']}/content"
        self.assertEqual(self.client.patch(
            path, json={"content": "<p>x</p>"},
        ).status_code, 401)
        self.assertEqual(self.edit(item, self.outsider).status_code, 404)
        self.assertEqual(self.edit(item, self.member).status_code, 403)
        self.assertEqual(self.edit({"id": self.material_set.id, "name": "set"}).status_code, 404)
        self.assertEqual(self.edit(item, content="<p><br></p>").status_code, 400)
        self.assertEqual(self.edit(item, title="Another title").status_code, 422)
        self.assertEqual(self.edit(item, name="Another name").status_code, 422)
        self.assertEqual(self.edit(item, content="x" * (1024 * 1024 + 1)).status_code, 422)
        other = self.create_copy(title="Existing").json()["data"]
        self.assertEqual(self.client.put(
            self.base + f"/materials/{item['id']}", headers=self.headers(),
            json={"name": other["name"]},
        ).status_code, 400)
        self.assertEqual(self.preview(item).json()["data"], original)
        renamed = self.client.put(
            self.base + f"/materials/{item['id']}", headers=self.headers(),
            json={"name": "Dotted.v2.md"},
        )
        self.assertEqual(renamed.status_code, 200)
        edited = self.edit(item, self.admin, content='<p onclick="bad()">New<script>bad</script></p>')
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertEqual(edited.json()["data"]["name"], "Dotted.v2.md")
        self.assertEqual(edited.json()["data"]["file_size"], len("<p>New</p>"))
        self.assertEqual(self.preview(item).json()["data"]["content"], "<p>New</p>")
        photo = self.upload("photo.png").json()["data"]
        self.assertEqual(self.edit(photo).status_code, 400)
        member_item = self.client.post(
            self.base + "/materials/copy", headers=self.headers(self.member),
            json={"material_set_id": self.material_set.id, "title": "Member", "content": "<p>Own</p>"},
        ).json()["data"]
        self.assertEqual(self.edit(member_item, self.member).status_code, 200)

    def test_legacy_documents_convert_on_read_and_save_without_original_loss(self):
        sources = (
            (".txt", b"Hello\nworld"), (".md", b"# Title"),
            (".markdown", b"**bold**"), (".pdf", self.pdf_bytes()),
            (".docx", self.docx_bytes()), (".html", b'<p onclick="bad()">Safe</p>'),
        )
        for extension, data in sources:
            with self.subTest(extension=extension):
                item = self.legacy(extension, data)
                response = self.preview(item)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["data"]["format"], "html")
                self.assertNotIn("onclick", response.json()["data"]["content"])
                stored = project_materials.get_project_material(self.owner["id"], self.project.id, item["id"])
                self.assertIsNone(stored.content_html)
                self.assertEqual(media_storage.read_media_bytes(item["object_key"]), data)
                edited = self.edit(item)
                self.assertEqual(edited.status_code, 200, edited.text)
                self.assertEqual(edited.json()["data"]["name"], item["name"])
                self.assertEqual(edited.json()["data"]["file_url"], "")
                storage.init_db()
                self.assertEqual(self.preview(item).json()["data"]["content"], "<p>Edited</p>")
                self.assertEqual(media_storage.read_media_bytes(item["object_key"]), data)
        legacy_alias = self.legacy("", b"# Alias", mime="text/x-markdown")
        self.assertIn("<h2>Alias</h2>", self.preview(legacy_alias).json()["data"]["content"])

    def test_migration_adds_nullable_content_without_reading_or_removing_originals(self):
        item = self.legacy(".html", b"<p>Old rich content</p>")
        with sqlite3.connect(self.db_path) as conn:
            conn.execute("ALTER TABLE project_materials DROP COLUMN content_html")
        with patch.object(media_storage, "read_media_bytes") as read, patch.object(media_storage, "delete_media") as delete:
            storage.init_db()
            storage.init_db()
            read.assert_not_called()
            delete.assert_not_called()
        self.assertEqual(media_storage.read_media_bytes(item["object_key"]), b"<p>Old rich content</p>")
        self.assertEqual(self.preview(item).json()["data"]["content"], "<p>Old rich content</p>")
        self.assertIsNone(project_materials.get_project_material(
            self.owner["id"], self.project.id, item["id"],
        ).content_html)

    def test_corrupt_empty_scanned_binary_and_oversize_documents_are_explicit_errors(self):
        for filename, data in (
            ("bad.pdf", b"not a PDF"), ("bad.docx", b"not a ZIP"),
            ("scan.pdf", self.pdf_bytes("")), ("empty.md", b"   "),
            ("binary.md", b"PK\x00\x01"), ("binary.txt", b"%PDF-\x00"),
            ("oversize.txt", b"x" * (1024 * 1024 + 1)),
        ):
            with self.subTest(filename=filename), patch.object(publishing, "put_media_bytes") as put:
                response = self.upload(filename, data)
                self.assertEqual(response.status_code, 400, response.text)
                self.assertTrue(response.json()["detail"])
                put.assert_not_called()
        with patch.object(document_copy.subprocess, "run", side_effect=subprocess.TimeoutExpired("parser", 15)):
            response = self.upload("timeout.pdf", self.pdf_bytes())
            self.assertEqual(response.status_code, 400)
            self.assertIn("time limit", response.json()["detail"])
        self.assertEqual(project_materials.list_project_materials(
            self.owner["id"], self.project.id, self.material_set.id,
        ), [])

    def test_parser_limits_and_docx_structure(self):
        with (
            patch.object(document_copy, "MAX_PDF_PAGES", 1),
            self.assertRaisesRegex(ValueError, "pages"),
        ):
            document_copy._convert_document(self.pdf_bytes(pages=2), ".pdf")
        with (
            patch.object(document_copy, "MAX_DOCX_EXPANDED_BYTES", 100),
            self.assertRaisesRegex(ValueError, "Expanded DOCX"),
        ):
            document_copy._convert_document(self.docx_bytes(), ".docx")
        with (
            patch.object(document_copy, "MAX_DOCX_ENTRIES", 1),
            self.assertRaisesRegex(ValueError, "archive entries"),
        ):
            document_copy._convert_document(self.docx_bytes(), ".docx")
        html = document_copy.parse_document_copy(self.docx_bytes(), ".docx")
        self.assertIn("<h2>Document title</h2>", html)
        self.assertIn("<ul><li>List item</li></ul>", html)
        self.assertIn("<p>Table cell</p>", html)
        with (
            patch.object(document_copy, "MAX_COPY_CONTENT_BYTES", 10),
            self.assertRaisesRegex(ValueError, "too large"),
        ):
            document_copy._convert_document(self.pdf_bytes(), ".pdf")
        with patch.object(publishing, "MAX_UPLOAD_SIZE_BYTES", 4):
            item = self.legacy(".txt", b"12345")
            self.assertEqual(self.preview(item).status_code, 400)

    def test_parser_preserves_markdown_code_and_never_loads_external_images(self):
        html = document_copy.parse_document_copy(
            b'```txt\nline 1\nline 2\n```\n\n![Alt text](https://example.com/private)\n\n<script>bad</script>',
            ".md",
        )
        self.assertIn("<p>line 1<br>line 2<br></p>", html)
        self.assertIn("Alt text", html)
        self.assertNotIn("<img", html)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        with self.assertRaisesRegex(ValueError, "nesting"):
            material_copy.sanitize_copy_html("<p>" * 129 + "text")
        with self.assertRaisesRegex(ValueError, "control characters"):
            material_copy.sanitize_copy_html("<p>\x00text</p>")

    def test_markdown_table_cells_and_link_labels_survive_editor_compatible_html(self):
        html = document_copy.parse_document_copy(
            b"| Heading | Other |\n| --- | --- |\n| **Cell** | [Visible label](https://example.com) |\n\n~~Removed~~",
            ".md",
        )
        self.assertIn("<p>Heading</p>", html)
        self.assertIn("<p><strong>Cell</strong></p>", html)
        self.assertIn("<p>Visible label</p>", html)
        self.assertIn("<s>Removed</s>", html)
        self.assertNotIn("<table", html)
        self.assertNotIn("<a ", html)
        self.assertNotIn("https://", html)
        response = self.upload("image-only.md", b"![](https://example.com/image.png)")
        self.assertEqual(response.status_code, 400, response.text)

    def test_docx_link_text_emphasis_and_consecutive_lists_are_preserved(self):
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = Document()
        paragraph = document.add_paragraph()
        link = OxmlElement("w:hyperlink")
        relation = paragraph.part.relate_to(
            "https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True,
        )
        link.set(qn("r:id"), relation)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        properties.append(OxmlElement("w:b"))
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = "Visible link text"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document.add_paragraph("First", "List Number")
        document.add_paragraph("Second", "List Number")
        buffer = io.BytesIO()
        document.save(buffer)
        html = document_copy.parse_document_copy(buffer.getvalue(), ".docx")
        self.assertIn("<strong>Visible link text</strong>", html)
        self.assertIn("<ol><li>First</li><li>Second</li></ol>", html)
        self.assertNotIn("<a ", html)
        self.assertNotIn("https://", html)

    def test_real_parser_timeout_and_busy_limits_are_explicit(self):
        start = time.monotonic()
        with (
            patch.object(document_copy, "PARSE_TIMEOUT_SECONDS", 0.001),
            self.assertRaisesRegex(ValueError, "time limit"),
        ):
            document_copy.parse_document_copy(self.pdf_bytes(), ".pdf")
        self.assertLess(time.monotonic() - start, 2)
        with (
            patch.object(document_copy, "_PARSER_SLOTS") as slots,
            self.assertRaisesRegex(ValueError, "busy"),
        ):
            slots.acquire.return_value = False
            document_copy.parse_document_copy(b"text", ".txt")
        slots.release.assert_not_called()
        self.assertEqual(document_copy.parse_document_copy(b"after timeout", ".txt"), "<p>after timeout</p>")

    def test_encrypted_pdf_empty_docx_and_page_limit_are_rejected(self):
        with fitz.open() as document:
            document.new_page().insert_text((72, 72), "Secret")
            encrypted = document.tobytes(
                encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="password", owner_pw="owner",
            )
        response = self.upload("encrypted.pdf", encrypted)
        self.assertEqual(response.status_code, 400)
        self.assertIn("Password-protected", response.json()["detail"])
        empty_docx = io.BytesIO()
        Document().save(empty_docx)
        response = self.upload("empty.docx", empty_docx.getvalue())
        self.assertEqual(response.status_code, 400)
        self.assertIn("empty", response.json()["detail"])
        response = self.upload("many-pages.pdf", self.pdf_bytes(pages=201))
        self.assertEqual(response.status_code, 400)
        self.assertIn("200 pages", response.json()["detail"])

    def test_extension_migration_is_once_only_and_skips_rich_titles(self):
        image = self.upload("launch.v1.png").json()["data"]
        rich = self.create_copy(title="Release.v2.html").json()["data"]
        with sqlite3.connect(self.db_path) as conn:
            conn.execute("DELETE FROM publishing_migrations WHERE key = 'stable-material-names-v1'")
            conn.execute("UPDATE project_materials SET name = ? WHERE id = ?", ("launch.v1.png", image["id"]))
        storage.init_db()
        storage.init_db()
        for item, expected in ((image, "launch.v1"), (rich, "Release.v2.html")):
            self.assertEqual(
                project_materials.get_project_material(self.owner["id"], self.project.id, item["id"]).name,
                expected,
            )
        self.client.put(
            self.base + f"/materials/{image['id']}", headers=self.headers(), json={"name": "Edited.png"},
        )
        storage.init_db()
        self.assertEqual(
            project_materials.get_project_material(self.owner["id"], self.project.id, image["id"]).name,
            "Edited.png",
        )

    def test_postgres_copy_excludes_initialized_migration_metadata(self):
        output = io.StringIO()
        with (
            patch("sys.argv", [
                "migrate_sqlite_to_postgres", "--sqlite-path", self.db_path,
                "--database-url", "postgresql://unused", "--dry-run",
            ]),
            redirect_stdout(output),
        ):
            self.assertEqual(migrate_sqlite_to_postgres.main(), 0)
        self.assertIn("project_materials: 1", output.getvalue())
        self.assertNotIn("publishing_migrations:", output.getvalue())
